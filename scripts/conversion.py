import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

INPUT_JSONL = "/home/getalp/bouarabs/Data_Generation/datasets(variantes)/mistral_5shot_v5.jsonl"
OUTPUT_XLSX = "/home/getalp/bouarabs/Data_Generation/out/reports_for_review.xlsx"

COLUMNS = [
    "record_id",
    "civilite", "nom", "prenom", "birth_date", "side", "subtype", "stage", "context",
    "cr_text",
]

def main():
    in_path = Path(INPUT_JSONL)
    out_path = Path(OUTPUT_XLSX)

    wb = Workbook()
    ws = wb.active
    ws.title = "reports"

    # Header style
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for j, col in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=j, value=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align

    wrap = Alignment(vertical="top", wrap_text=True)

    row = 2
    with in_path.open("r", encoding="utf-8") as f:
        for i, line in enumerate(f, start=1):
            line = line.strip()
            if not line:
                continue
            obj = json.loads(line)

            fields = obj.get("fields") or {}
            cr_text = obj.get("cr_text", "")

            data = {
                "record_id": i,
                **fields,
                "cr_text": cr_text,
            }

            for j, col in enumerate(COLUMNS, start=1):
                ws.cell(row=row, column=j, value=data.get(col, ""))
                ws.cell(row=row, column=j).alignment = wrap
            row += 1

    # Column widths (adjust as needed)
    widths = {
        "record_id": 9,
        "civilite": 10, "nom": 14, "prenom": 14, "birth_date": 12,
        "side": 12, "subtype": 12, "stage": 12, "context": 14,
        "cr_text": 80,
    }
    for j, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(j)].width = widths.get(col, 15)

    ws.freeze_panes = "A2"
    wb.save(out_path)
    print(f"Saved: {out_path.resolve()}")

if __name__ == "__main__":
    main()
